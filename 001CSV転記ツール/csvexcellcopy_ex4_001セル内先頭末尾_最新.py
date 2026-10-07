import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import csv
import chardet
import os
import json
from datetime import datetime
import re
from uuid import uuid4
import openpyxl
from openpyxl.utils import get_column_letter

class CSVProcessorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("CSV/Excel Processor")
        self.file_configs = []
        self.file_data = []
        self.file_headers = []
        self.set_frames = []
        self.settings_file = "settings.json"
        self.current_set_name = "新規"
        self.load_settings()
        self.create_gui()

    def load_settings(self):
        if os.path.exists(self.settings_file):
            with open(self.settings_file, 'r', encoding='utf-8') as f:
                self.saved_settings = json.load(f)
        else:
            self.saved_settings = {}

    def save_settings(self):
        with open(self.settings_file, 'w', encoding='utf-8') as f:
            json.dump(self.saved_settings, f, ensure_ascii=False, indent=2)

    def create_gui(self):
        self.main_frame = ttk.Frame(self.root)
        self.main_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        # 設定名
        set_name_frame = ttk.Frame(self.main_frame)
        set_name_frame.grid(row=0, column=0, sticky=tk.W, pady=2)
        ttk.Label(set_name_frame, text="設定:").grid(row=0, column=0, sticky=tk.W, padx=(0, 5))
        self.set_name_combo = ttk.Combobox(set_name_frame, values=["新規"] + list(self.saved_settings.keys()), width=20)
        self.set_name_combo.set(self.current_set_name)
        self.set_name_combo.grid(row=0, column=1, sticky=tk.W, padx=(0, 5))
        self.set_name_combo.bind('<<ComboboxSelected>>', self.load_selected_set)

        # ファイル数とセット数
        count_frame = ttk.Frame(self.main_frame)
        count_frame.grid(row=1, column=0, sticky=tk.W, pady=2)
        ttk.Label(count_frame, text="ファイル数:").grid(row=0, column=0, sticky=tk.W, padx=(0, 5))
        self.file_count_var = tk.StringVar(value="2")
        self.file_count_entry = ttk.Entry(count_frame, textvariable=self.file_count_var, width=5)
        self.file_count_entry.grid(row=0, column=1, sticky=tk.W, padx=(0, 5))
        self.file_count_entry.bind("<MouseWheel>", self.adjust_file_count)
        ttk.Button(count_frame, text="確定", command=self.update_file_configs).grid(row=0, column=2, sticky=tk.W, padx=(0, 10))
        ttk.Label(count_frame, text="セット数:").grid(row=0, column=3, sticky=tk.W, padx=(0, 5))
        self.set_count_var = tk.StringVar(value="1")
        self.set_count_entry = ttk.Entry(count_frame, textvariable=self.set_count_var, width=5)
        self.set_count_entry.grid(row=0, column=4, sticky=tk.W, padx=(0, 5))
        self.set_count_entry.bind("<MouseWheel>", self.adjust_set_count)
        ttk.Button(count_frame, text="確定", command=self.update_sets).grid(row=0, column=5, sticky=tk.W, padx=(0, 5))

        # ファイル設定フレーム
        self.file_frame = ttk.Frame(self.main_frame)
        self.file_frame.grid(row=2, column=0, sticky=tk.W+tk.E, pady=2)
        self.update_file_configs()

        # セットフレーム
        self.sets_frame = ttk.Frame(self.main_frame)
        self.sets_frame.grid(row=3, column=0, sticky=tk.W+tk.E, pady=2)
        self.update_sets()

        # 保存・削除・実行ボタン
        button_frame = ttk.Frame(self.main_frame)
        button_frame.grid(row=4, column=0, sticky=tk.W, pady=5)
        ttk.Button(button_frame, text="設定保存", command=self.save_current_setting).grid(row=0, column=0, padx=(0, 5))
        ttk.Button(button_frame, text="設定削除", command=self.delete_current_setting).grid(row=0, column=1, padx=(0, 5))
        ttk.Button(button_frame, text="実行", command=self.execute, style='Large.TButton').grid(row=0, column=2, padx=(0, 5))

        # スタイル設定
        style = ttk.Style()
        style.configure('Large.TButton', font=('Helvetica', 12), padding=10)

    def adjust_file_count(self, event):
        current = int(self.file_count_var.get())
        if event.delta > 0:
            self.file_count_var.set(str(current + 1))
        elif current > 1:
            self.file_count_var.set(str(current - 1))
        self.update_file_configs()

    def adjust_set_count(self, event):
        current = int(self.set_count_var.get())
        if event.delta > 0:
            self.set_count_var.set(str(current + 1))
        elif current > 1:
            self.set_count_var.set(str(current - 1))
        self.update_sets()

    def detect_encoding(self, file_path):
        with open(file_path, 'rb') as f:
            raw_data = f.read(10000)
            result = chardet.detect(raw_data)
            return result['encoding'] or 'utf-8'

    def detect_delimiter(self, file_path, encoding):
        with open(file_path, 'r', encoding=encoding) as f:
            first_line = f.readline()
            possible_delimiters = [',', ';', '\t', '|']
            return max(possible_delimiters, key=lambda d: first_line.count(d))

    def load_file_headers(self, file_path, encoding, has_header, sheet_name=None, read_format="表示値"):
        if not file_path:
            return [f"列{j+1}" for j in range(10)]
        ext = os.path.splitext(file_path)[1].lower()
        if ext in ['.xlsx', '.xls']:
            try:
                wb = openpyxl.load_workbook(file_path, data_only=(read_format == "内部値"))
                sheet = wb[sheet_name] if sheet_name else wb.worksheets[0]
                first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
                if not first_row or all(cell is None or str(cell).strip() == "" for cell in first_row):
                    wb.close()
                    return [f"列{j+1}" for j in range(10)]
                headers = [str(cell) if cell is not None else "" for cell in first_row] if has_header == "カラム有" else [f"列{j+1}" for j in range(len(first_row))]
                wb.close()
                return headers
            except Exception:
                return [f"列{j+1}" for j in range(10)]
        else:  # CSV
            encoding = encoding if encoding != "自動" else self.detect_encoding(file_path)
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    first_line = f.readline().strip()
                    if not first_line:
                        return [f"列{j+1}" for j in range(10)]
                    delimiter = self.detect_delimiter(file_path, encoding)
                    f.seek(0)
                    reader = csv.reader(f, delimiter=delimiter)
                    first_row = next(reader, None)
                    if not first_row or all(not cell.strip() for cell in first_row):
                        return [f"列{j+1}" for j in range(10)]
                    return first_row if has_header == "カラム有" else [f"列{j+1}" for j in range(len(first_row))]
            except Exception:
                return [f"列{j+1}" for j in range(10)]

    def get_excel_sheets(self, file_path):
        if os.path.splitext(file_path)[1].lower() in ['.xlsx', '.xls']:
            try:
                wb = openpyxl.load_workbook(file_path)
                sheet_names = wb.sheetnames
                wb.close()
                return sheet_names
            except Exception:
                return []
        return []

    def update_file_configs(self):
        current_configs = [
            {
                "path": c["path"].get(),
                "has_header": c["has_header"].get(),
                "encoding": c["encoding"].get(),
                "sheet": c["sheet"].get(),
                "memo": c["memo"].get(),
                "headers": c["headers"],
                "read_format": c["read_format"].get() if c.get("read_format") else "表示値"
            } for c in self.file_configs
        ]
        
        for widget in self.file_frame.winfo_children():
            widget.destroy()
        
        new_count = int(self.file_count_var.get())
        self.file_configs = []
        encodings = ["自動", "shift-jis", "utf-8", "iso-2022-jp", "euc-jp", "CP932"]
        read_formats = ["表示値", "内部値"]
        
        for i in range(new_count):
            frame = ttk.Frame(self.file_frame)
            frame.grid(row=i, column=0, sticky=tk.W, pady=2)
            ttk.Label(frame, text=f"ファイル{i+1}:").grid(row=0, column=0, sticky=tk.W, padx=(0, 5))
            
            if i < len(current_configs) and self.current_set_name != "新規":
                config = current_configs[i]
                path_var = tk.StringVar(value=config["path"])
                has_header_var = tk.StringVar(value=config["has_header"])
                encoding_var = tk.StringVar(value=config["encoding"])
                sheet_var = tk.StringVar(value=config["sheet"])
                memo_var = tk.StringVar(value=config["memo"])
                read_format_var = tk.StringVar(value=config["read_format"])
                headers = config["headers"]
            else:
                path_var = tk.StringVar()
                has_header_var = tk.StringVar(value="カラム有")
                encoding_var = tk.StringVar(value="自動")
                sheet_var = tk.StringVar(value="")
                memo_var = tk.StringVar(value="")
                read_format_var = tk.StringVar(value="表示値")
                headers = [f"列{j+1}" for j in range(10)]
            
            path_entry = ttk.Entry(frame, textvariable=path_var, width=45)
            path_entry.grid(row=0, column=1, sticky=tk.W, padx=(0, 5))
            ttk.Button(frame, text="参照", command=lambda v=path_var, idx=i: self.browse_file(v, idx)).grid(row=0, column=2, sticky=tk.W, padx=(0, 5))
            header_combo = ttk.Combobox(frame, values=["カラム有", "カラム無"], textvariable=has_header_var, width=8)
            header_combo.grid(row=0, column=3, sticky=tk.W, padx=(0, 5))
            ttk.Label(frame, text="文字コード:").grid(row=0, column=4, sticky=tk.W, padx=(0, 5))
            encoding_combo = ttk.Combobox(frame, values=encodings, textvariable=encoding_var, width=10)
            encoding_combo.grid(row=0, column=5, sticky=tk.W, padx=(0, 5))
            sheet_combo = ttk.Combobox(frame, textvariable=sheet_var, width=15)
            sheet_combo.grid(row=0, column=6, sticky=tk.W, padx=(0, 5))
            sheet_combo.grid_remove()
            ttk.Label(frame, text="読込:").grid(row=0, column=7, sticky=tk.W, padx=(0, 5))
            read_format_combo = ttk.Combobox(frame, values=read_formats, textvariable=read_format_var, width=10)
            read_format_combo.grid(row=0, column=8, sticky=tk.W, padx=(0, 5))
            read_format_combo.grid_remove()
            ttk.Button(frame, text="クリア", command=lambda v=path_var, idx=i: self.clear_file(v, idx)).grid(row=0, column=9, sticky=tk.W, padx=(0, 5))
            ttk.Label(frame, text="メモ:").grid(row=0, column=10, sticky=tk.W, padx=(0, 5))
            memo_entry = ttk.Entry(frame, textvariable=memo_var, width=20)
            memo_entry.grid(row=0, column=11, sticky=tk.W, padx=(0, 5))
            
            self.file_configs.append({
                "path": path_var,
                "has_header": has_header_var,
                "encoding": encoding_var,
                "sheet": sheet_var,
                "memo": memo_var,
                "headers": headers,
                "sheet_combo": sheet_combo,
                "read_format": read_format_var,
                "read_format_combo": read_format_combo
            })
            
            if path_var.get() and os.path.splitext(path_var.get())[1].lower() in ['.xlsx', '.xls']:
                sheet_names = self.get_excel_sheets(path_var.get())
                sheet_combo.configure(values=sheet_names)
                if sheet_var.get() in sheet_names:
                    sheet_combo.grid()
                    read_format_combo.grid()
                else:
                    sheet_var.set(sheet_names[0] if sheet_names else "")
                    sheet_combo.grid() if sheet_names else sheet_combo.grid_remove()
                    read_format_combo.grid() if sheet_names else read_format_combo.grid_remove()
            else:
                sheet_var.set("")
                sheet_combo.grid_remove()
                read_format_combo.grid_remove()
            
            path_var.trace_add("write", lambda *args, v=path_var, idx=i: self.update_file_headers(v, idx))
            has_header_var.trace_add("write", lambda *args, v=path_var, idx=i: self.update_file_headers(v, idx))
            encoding_var.trace_add("write", lambda *args, v=path_var, idx=i: self.update_file_headers(v, idx))
            sheet_var.trace_add("write", lambda *args, v=path_var, idx=i: self.update_file_headers(v, idx))
            read_format_var.trace_add("write", lambda *args, v=path_var, idx=i: self.update_file_headers(v, idx))
        
        self.update_all_column_menus()

    def browse_file(self, path_var, idx):
        file_path = filedialog.askopenfilename(filetypes=[("CSV/Excel files", "*.csv;*.xlsx;*.xls")])
        if file_path:
            path_var.set(file_path)
            self.update_file_headers(path_var, idx)
            ext = os.path.splitext(file_path)[1].lower()
            if ext in ['.xlsx', '.xls']:
                sheet_names = self.get_excel_sheets(file_path)
                self.file_configs[idx]["sheet_combo"].configure(values=sheet_names)
                self.file_configs[idx]["sheet"].set(sheet_names[0] if sheet_names else "")
                self.file_configs[idx]["sheet_combo"].grid()
                self.file_configs[idx]["read_format_combo"].grid()
            else:
                self.file_configs[idx]["sheet_combo"].grid_remove()
                self.file_configs[idx]["read_format_combo"].grid_remove()
                self.file_configs[idx]["sheet"].set("")
            self.update_all_column_menus()

    def clear_file(self, path_var, idx):
        path_var.set("")
        self.file_configs[idx]["headers"] = [f"列{j+1}" for j in range(10)]
        self.file_configs[idx]["sheet"].set("")
        self.file_configs[idx]["memo"].set("")
        self.file_configs[idx]["read_format"].set("表示値")
        self.file_configs[idx]["sheet_combo"].grid_remove()
        self.file_configs[idx]["read_format_combo"].grid_remove()
        self.update_all_column_menus()

    def update_file_headers(self, path_var, idx):
        sheet_name = self.file_configs[idx]["sheet"].get() if self.file_configs[idx]["sheet"].get() else None
        headers = self.load_file_headers(
            path_var.get(),
            self.file_configs[idx]["encoding"].get(),
            self.file_configs[idx]["has_header"].get(),
            sheet_name,
            self.file_configs[idx]["read_format"].get()
        )
        if headers:
            self.file_configs[idx]["headers"] = headers
        self.update_all_column_menus()

    def update_all_column_menus(self):
        if not hasattr(self, 'set_frames') or not self.set_frames:
            return
        file_options = [f"ファイル{j+1}" for j in range(len(self.file_configs))]
        date_formats = [
            "%Y/%m/%d", "%Y年%m月%d日", "%Y%m%d", "%Y-%m-%d",
            "%m/%d/%Y", "%d-%m-%Y", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S"
        ]
        date_options = [datetime.now().strftime(fmt) for fmt in date_formats]
        for set_frame in self.set_frames:
            condition_frame = set_frame.winfo_children()[0]
            for cond in condition_frame.conditions:
                current_csv1 = cond["csv1"].get()
                current_csv2 = cond["csv2"].get()
                cond["csv1"].configure(values=file_options)
                cond["csv2"].configure(values=file_options + ["入力文字"])
                cond["csv1"].set(current_csv1 if current_csv1 in file_options else file_options[0] if file_options else "")
                cond["csv2"].set(current_csv2 if current_csv2 in file_options + ["入力文字"] else file_options[1] if len(file_options) > 1 else file_options[0] if file_options else "入力文字")
                
                csv1_idx = int(cond["csv1"].get().replace("ファイル", "")) - 1 if cond["csv1"].get().startswith("ファイル") else 0
                if csv1_idx < len(self.file_configs):
                    headers = self.file_configs[csv1_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                    cond["column1"].configure(values=headers)
                    current_column1 = cond["column1"].get()
                    cond["column1"].set(current_column1 if current_column1 in headers else headers[0] if headers else "")
                else:
                    cond["column1"].configure(values=[f"列{j+1}" for j in range(10)])
                    cond["column1"].set("列1")
                
                csv2_value = cond["csv2"].get()
                if csv2_value != "入力文字" and csv2_value.startswith("ファイル"):
                    csv2_idx = int(csv2_value.replace("ファイル", "")) - 1
                    if csv2_idx < len(self.file_configs):
                        headers = self.file_configs[csv2_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                        cond["column2"].configure(values=headers)
                        current_column2 = cond["column2"].get()
                        cond["column2"].set(current_column2 if current_column2 in headers else headers[0] if headers else "")
                    else:
                        cond["column2"].configure(values=[f"列{j+1}" for j in range(10)])
                        cond["column2"].set("列1")
                else:
                    cond["column2"].configure(values=[])
                    cond["column2"].set("")
                
                self.update_column2(cond["csv2"], cond["value"], cond["column2"], cond["no_label"], cond["to_label"])
                
                cond["csv1"].bind("<<ComboboxSelected>>", lambda e, c1=cond["csv1"], col1=cond["column1"]: col1.configure(values=self.file_configs[int(c1.get().replace("ファイル", ""))-1]["headers"] or [f"列{j+1}" for j in range(10)] if c1.get().startswith("ファイル") and int(c1.get().replace("ファイル", ""))-1 < len(self.file_configs) else [f"列{j+1}" for j in range(10)]))
                cond["csv2"].bind("<<ComboboxSelected>>", lambda e, c2=cond["csv2"], v=cond["value"], c1=cond["column2"], nl=cond["no_label"], tl=cond["to_label"]: self.update_column2(c2, v, c1, nl, tl))
            
            action_frame = set_frame.winfo_children()[1]
            for action in action_frame.actions:
                current_source = action["source"].get()
                current_dest = action["dest"].get()
                action["source"].configure(values=file_options + ["本日の日時", "入力文字"])
                action["dest"].configure(values=file_options)
                action["source"].set(current_source if current_source in file_options + ["本日の日時", "入力文字"] else file_options[0] if file_options else "入力文字")
                action["dest"].set(current_dest if current_dest in file_options else file_options[1] if len(file_options) > 1 else file_options[0] if file_options else "")
                
                source_value = action["source"].get()
                if source_value == "本日の日時":
                    action["source_column"].configure(values=date_options)
                    current_source_column = action["source_column"].get()
                    action["source_column"].set(current_source_column if current_source_column in date_options else date_options[0])
                    action["source_column"].grid(row=0, column=3, padx=(0, 5))
                    action["source_column_label"].grid_forget()
                    action["value"].grid_forget()
                elif source_value.startswith("ファイル"):
                    source_idx = int(source_value.replace("ファイル", "")) - 1
                    if source_idx < len(self.file_configs):
                        headers = self.file_configs[source_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                        action["source_column"].configure(values=headers)
                        current_source_column = action["source_column"].get()
                        action["source_column"].set(current_source_column if current_source_column in headers else headers[0] if headers else "")
                        action["source_column"].grid(row=0, column=3, padx=(0, 5))
                        action["source_column_label"].grid_forget()
                    else:
                        action["source_column"].configure(values=[f"列{j+1}" for j in range(10)])
                        action["source_column"].set("列1")
                        action["source_column"].grid(row=0, column=3, padx=(0, 5))
                        action["source_column_label"].grid_forget()
                    action["value"].grid_forget()
                else:
                    action["source_column"].grid_forget()
                    action["source_column_label"].grid_forget()
                    action["value"].grid(row=0, column=3, padx=(0, 5))
                
                dest_idx = int(action["dest"].get().replace("ファイル", "")) - 1 if action["dest"].get().startswith("ファイル") else 0
                if dest_idx < len(self.file_configs):
                    headers = self.file_configs[dest_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                    action["dest_column"].configure(values=headers)
                    current_dest_column = action["dest_column"].get()
                    action["dest_column"].set(current_dest_column if current_dest_column in headers else headers[0] if headers else "")
                else:
                    action["dest_column"].configure(values=[f"列{j+1}" for j in range(10)])
                    action["dest_column"].set("列1")
                
                ext = os.path.splitext(self.file_configs[dest_idx]["path"].get())[1].lower() if dest_idx < len(self.file_configs) else ""
                if ext in ['.xlsx', '.xls']:
                    action["write_format"].master.grid(row=0, column=12, padx=(0, 5))
                    action["write_format_label"].grid(row=0, column=11, padx=(0, 5))
                else:
                    action["write_format"].master.grid_remove()
                    action["write_format_label"].grid_remove()
                
                action["source"].bind("<<ComboboxSelected>>", lambda e, s=action["source"], sc=action["source_column"], scl=action["source_column_label"], v=action["value"]: self.update_source_column(s, sc, scl, v))
                action["dest"].bind("<<ComboboxSelected>>", lambda e, d=action["dest"], dc=action["dest_column"], wf=action["write_format"], wfl=action["write_format_label"]: self.update_dest_column(d, dc, wf, wfl))

    def update_sets(self):
        current_sets = []
        for set_frame in self.set_frames:
            condition_frame = set_frame.winfo_children()[0]
            action_frame = set_frame.winfo_children()[1]
            conditions = [
                {
                    "csv1": c["csv1"].get(),
                    "column1": c["column1"].get(),
                    "csv2": c["csv2"].get(),
                    "column2": c["column2"].get(),
                    "operator": c["operator"].get(),
                    "value": c["value"].get(),
                    "logic": c["logic"].get() if c["logic"] else ""
                } for c in condition_frame.conditions
            ]
            actions = [
                {
                    "source": a["source"].get(),
                    "source_column": a["source_column"].get(),
                    "dest": a["dest"].get(),
                    "dest_column": a["dest_column"].get(),
                    "method": a["method"].get(),
                    "value": a["value"].get(),
                    "write_format": a["write_format"].get()
                } for a in action_frame.actions
            ]
            current_sets.append({"conditions": conditions, "actions": actions})
        
        for frame in self.set_frames:
            frame.destroy()
        self.set_frames = []
        
        count = int(self.set_count_var.get())
        for i in range(count):
            set_frame = ttk.LabelFrame(self.sets_frame, text=f"セット{i+1}")
            set_frame.grid(row=i, column=0, sticky=tk.W+tk.E, pady=2)
            self.set_frames.append(set_frame)
            condition_count_var = self.create_condition_frame(set_frame, i)
            action_count_var = self.create_action_frame(set_frame, i)
            
            if self.current_set_name != "新規" and i < len(current_sets):
                condition_frame = set_frame.winfo_children()[0]
                condition_count_var.set(len(current_sets[i]["conditions"]) or 1)
                self.update_conditions(i, condition_count_var)
                for j, cond in enumerate(current_sets[i]["conditions"]):
                    if j < len(condition_frame.conditions):
                        condition_frame.conditions[j]["csv1"].set(cond["csv1"])
                        condition_frame.conditions[j]["column1"].set(cond["column1"])
                        condition_frame.conditions[j]["csv2"].set(cond["csv2"])
                        condition_frame.conditions[j]["column2"].set(cond["column2"])
                        condition_frame.conditions[j]["operator"].set(cond["operator"])
                        condition_frame.conditions[j]["value"].delete(0, tk.END)
                        condition_frame.conditions[j]["value"].insert(0, cond["value"])
                        if condition_frame.conditions[j]["logic"]:
                            condition_frame.conditions[j]["logic"].set(cond["logic"])
                        self.update_column2(
                            condition_frame.conditions[j]["csv2"],
                            condition_frame.conditions[j]["value"],
                            condition_frame.conditions[j]["column2"],
                            condition_frame.conditions[j]["no_label"],
                            condition_frame.conditions[j]["to_label"]
                        )
                
                action_frame = set_frame.winfo_children()[1]
                action_count_var.set(len(current_sets[i]["actions"]) or 1)
                self.update_actions(i, action_count_var)
                for j, act in enumerate(current_sets[i]["actions"]):
                    if j < len(action_frame.actions):
                        action_frame.actions[j]["source"].set(act["source"])
                        action_frame.actions[j]["source_column"].set(act["source_column"])
                        action_frame.actions[j]["dest"].set(act["dest"])
                        action_frame.actions[j]["dest_column"].set(act["dest_column"])
                        action_frame.actions[j]["method"].set(act["method"])
                        action_frame.actions[j]["value"].delete(0, tk.END)
                        action_frame.actions[j]["value"].insert(0, act["value"])
                        action_frame.actions[j]["write_format"].set(act["write_format"])
                        self.update_source_column(
                            action_frame.actions[j]["source"],
                            action_frame.actions[j]["source_column"],
                            action_frame.actions[j]["source_column_label"],
                            action_frame.actions[j]["value"]
                        )
                        ext = os.path.splitext(self.file_configs[int(act["dest"].replace("ファイル", ""))-1]["path"].get())[1].lower() if act["dest"].startswith("ファイル") else ""
                        if ext in ['.xlsx', '.xls']:
                            action_frame.actions[j]["write_format"].master.grid(row=0, column=12, padx=(0, 5))
                            action_frame.actions[j]["write_format_label"].grid(row=0, column=11, padx=(0, 5))
                        else:
                            action_frame.actions[j]["write_format"].master.grid_remove()
                            action_frame.actions[j]["write_format_label"].grid_remove()
            else:
                condition_count_var.set("1")
                action_count_var.set("1")
                self.update_conditions(i, condition_count_var)
                self.update_actions(i, action_count_var)

    def create_condition_frame(self, parent, set_idx):
        condition_frame = ttk.Frame(parent)
        condition_frame.pack(fill=tk.X, anchor=tk.NW, pady=2)
        count_subframe = ttk.Frame(condition_frame)
        count_subframe.pack(side=tk.LEFT, anchor=tk.NW)
        ttk.Label(count_subframe, text="条件数:").pack(side=tk.LEFT, padx=(0, 5))
        condition_count_var = tk.StringVar(value="1")
        entry = ttk.Entry(count_subframe, textvariable=condition_count_var, width=5)
        entry.pack(side=tk.LEFT, padx=(0, 5))
        entry.bind("<MouseWheel>", lambda e: self.adjust_condition_count(e, set_idx, condition_count_var))
        ttk.Button(count_subframe, text="確定", command=lambda: self.update_conditions(set_idx, condition_count_var)).pack(side=tk.LEFT, padx=(0, 5))
        condition_frame.conditions_subframe = ttk.Frame(condition_frame)
        condition_frame.conditions_subframe.pack(side=tk.LEFT, anchor=tk.NW, padx=(20, 0))
        condition_frame.conditions = []
        condition_frame.condition_count_var = condition_count_var
        self.update_conditions(set_idx, condition_count_var)
        return condition_count_var

    def adjust_condition_count(self, event, set_idx, condition_count_var):
        current = int(condition_count_var.get())
        if event.delta > 0:
            condition_count_var.set(str(current + 1))
        elif current > 1:
            condition_count_var.set(str(current - 1))
        self.update_conditions(set_idx, condition_count_var)

    def update_conditions(self, set_idx, condition_count_var):
        condition_frame = self.set_frames[set_idx].winfo_children()[0]
        conditions_subframe = condition_frame.conditions_subframe
        
        current_conditions = [
            {
                "csv1": c["csv1"].get(),
                "column1": c["column1"].get(),
                "csv2": c["csv2"].get(),
                "column2": c["column2"].get(),
                "operator": c["operator"].get(),
                "value": c["value"].get(),
                "logic": c["logic"].get() if c["logic"] else ""
            } for c in condition_frame.conditions
        ] if self.current_set_name != "新規" else []
        
        for widget in conditions_subframe.winfo_children():
            widget.destroy()
        condition_frame.conditions = []
        
        count = max(int(condition_count_var.get()), 1)
        file_options = [f"ファイル{j+1}" for j in range(len(self.file_configs))]
        
        for i in range(count):
            frame = ttk.Frame(conditions_subframe)
            frame.pack(fill=tk.X, anchor=tk.NW, pady=2)
            ttk.Label(frame, text=f"条件{i+1}").pack(side=tk.LEFT, padx=(0, 5))
            csv1 = ttk.Combobox(frame, values=file_options, width=10)
            csv1.pack(side=tk.LEFT, padx=(0, 5))
            ttk.Label(frame, text="の").pack(side=tk.LEFT, padx=(0, 5))
            column1 = ttk.Combobox(frame, width=15)
            column1.pack(side=tk.LEFT, padx=(0, 5))
            ttk.Label(frame, text="が").pack(side=tk.LEFT, padx=(0, 5))
            csv2 = ttk.Combobox(frame, values=file_options + ["入力文字"], width=10)
            csv2.pack(side=tk.LEFT, padx=(0, 5))
            no_label = ttk.Label(frame, text="の")
            no_label.pack(side=tk.LEFT, padx=(0, 5))
            value_entry = ttk.Entry(frame, width=18)
            column2 = ttk.Combobox(frame, width=15)
            to_label = ttk.Label(frame, text="と")
            to_label.pack(side=tk.LEFT, padx=(0, 5))
            operator = ttk.Combobox(frame, values=["一致する", "異なる", "含む", "含まない", "より大きい", "より小さい"], width=10)
            operator.pack(side=tk.LEFT, padx=(0, 5))
            ttk.Label(frame, text="場合").pack(side=tk.LEFT, padx=(0, 5))
            logic = None
            if i < count - 1:
                logic = ttk.Combobox(frame, values=["かつ", "または"], width=5)
                logic.pack(side=tk.LEFT, padx=(0, 5))
            
            if self.current_set_name != "新規" and i < len(current_conditions):
                cond = current_conditions[i]
                csv1.set(cond["csv1"] if cond["csv1"] in file_options else file_options[0] if file_options else "")
                csv1_idx = int(csv1.get().replace("ファイル", "")) - 1 if csv1.get().startswith("ファイル") else 0
                headers = self.file_configs[csv1_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                column1.configure(values=headers)
                column1.set(cond["column1"] if cond["column1"] in headers else headers[0] if headers else "")
                csv2.set(cond["csv2"] if cond["csv2"] in file_options + ["入力文字"] else file_options[1] if len(file_options) > 1 else file_options[0] if file_options else "入力文字")
                if csv2.get() == "入力文字":
                    value_entry.pack(side=tk.LEFT, padx=(0, 5), after=no_label)
                    value_entry.delete(0, tk.END)
                    value_entry.insert(0, cond["value"])
                else:
                    csv2_idx = int(csv2.get().replace("ファイル", "")) - 1 if csv2.get().startswith("ファイル") else 0
                    headers2 = self.file_configs[csv2_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                    column2.configure(values=headers2)
                    column2.set(cond["column2"] if cond["column2"] in headers2 else headers2[0] if headers2 else "")
                    column2.pack(side=tk.LEFT, padx=(0, 5), after=no_label)
                operator.set(cond["operator"] if cond["operator"] in ["一致する", "異なる", "含む", "含まない", "より大きい", "より小さい"] else "一致する")
                if logic and cond["logic"]:
                    logic.set(cond["logic"] if cond["logic"] in ["かつ", "または"] else "かつ")
            else:
                csv1.set(file_options[0] if file_options else "")
                csv1_idx = int(csv1.get().replace("ファイル", "")) - 1 if csv1.get().startswith("ファイル") else 0
                headers = self.file_configs[csv1_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                column1.configure(values=headers)
                column1.set(headers[0] if headers else "")
                csv2.set(file_options[1] if len(file_options) > 1 else file_options[0] if file_options else "入力文字")
                if csv2.get() == "入力文字":
                    value_entry.pack(side=tk.LEFT, padx=(0, 5), after=no_label)
                else:
                    csv2_idx = int(csv2.get().replace("ファイル", "")) - 1 if csv2.get().startswith("ファイル") else 0
                    headers2 = self.file_configs[csv2_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                    column2.configure(values=headers2)
                    column2.set(headers2[0] if headers2 else "")
                    column2.pack(side=tk.LEFT, padx=(0, 5), after=no_label)
                operator.set("一致する")
                if logic:
                    logic.set("かつ")
            
            condition_frame.conditions.append({
                "csv1": csv1, "column1": column1, "csv2": csv2, "column2": column2,
                "operator": operator, "value": value_entry, "logic": logic,
                "no_label": no_label, "to_label": to_label
            })
            
            csv1.bind("<<ComboboxSelected>>", lambda e, c1=csv1, col1=column1: col1.configure(values=self.file_configs[int(c1.get().replace("ファイル", ""))-1]["headers"] or [f"列{j+1}" for j in range(10)] if c1.get().startswith("ファイル") and int(c1.get().replace("ファイル", ""))-1 < len(self.file_configs) else [f"列{j+1}" for j in range(10)]))
            csv2.bind("<<ComboboxSelected>>", lambda e, c2=csv2, v=value_entry, c1=column2, nl=no_label, tl=to_label: self.update_column2(c2, v, c1, nl, tl))
            
            self.update_column2(csv2, value_entry, column2, no_label, to_label)

    def update_column2(self, csv2, value_entry, column2, no_label, to_label):
        if csv2.get() == "入力文字":
            column2.pack_forget()
            no_label.pack(side=tk.LEFT, padx=(0, 5), after=csv2)
            value_entry.pack(side=tk.LEFT, padx=(0, 5), after=no_label)
            to_label.pack(side=tk.LEFT, padx=(0, 5), after=value_entry)
        else:
            value_entry.pack_forget()
            no_label.pack(side=tk.LEFT, padx=(0, 5), after=csv2)
            column2.pack(side=tk.LEFT, padx=(0, 5), after=no_label)
            to_label.pack(side=tk.LEFT, padx=(0, 5), after=column2)
            if csv2.get().startswith("ファイル"):
                csv2_idx = int(csv2.get().replace("ファイル", "")) - 1
                if csv2_idx < len(self.file_configs):
                    headers = self.file_configs[csv2_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                    column2.configure(values=headers)
                    current_column2 = column2.get()
                    column2.set(current_column2 if current_column2 in headers else headers[0] if headers else "")
                else:
                    column2.configure(values=[f"列{j+1}" for j in range(10)])
                    column2.set("列1")

    def create_action_frame(self, parent, set_idx):
        action_frame = ttk.Frame(parent)
        action_frame.pack(fill=tk.X, anchor=tk.NW, pady=2)
        count_subframe = ttk.Frame(action_frame)
        count_subframe.pack(side=tk.LEFT, anchor=tk.NW)
        ttk.Label(count_subframe, text="実行数:").pack(side=tk.LEFT, padx=(0, 5))
        action_count_var = tk.StringVar(value="1")
        entry = ttk.Entry(count_subframe, textvariable=action_count_var, width=5)
        entry.pack(side=tk.LEFT, padx=(0, 5))
        entry.bind("<MouseWheel>", lambda e: self.adjust_action_count(e, set_idx, action_count_var))
        ttk.Button(count_subframe, text="確定", command=lambda: self.update_actions(set_idx, action_count_var)).pack(side=tk.LEFT, padx=(0, 5))
        action_frame.actions_subframe = ttk.Frame(action_frame)
        action_frame.actions_subframe.pack(side=tk.LEFT, anchor=tk.NW, padx=(20, 0))
        action_frame.actions = []
        action_frame.action_count_var = action_count_var
        self.update_actions(set_idx, action_count_var)
        return action_count_var

    def adjust_action_count(self, event, set_idx, action_count_var):
        current = int(action_count_var.get())
        if event.delta > 0:
            action_count_var.set(str(current + 1))
        elif current > 1:
            action_count_var.set(str(current - 1))
        self.update_actions(set_idx, action_count_var)

    def update_actions(self, set_idx, action_count_var):
        action_frame = self.set_frames[set_idx].winfo_children()[1]
        actions_subframe = action_frame.actions_subframe
        
        current_actions = [
            {
                "source": a["source"].get(),
                "source_column": a["source_column"].get(),
                "dest": a["dest"].get(),
                "dest_column": a["dest_column"].get(),
                "method": a["method"].get(),
                "value": a["value"].get(),
                "write_format": a["write_format"].get()
            } for a in action_frame.actions
        ] if self.current_set_name != "新規" else []
        
        for widget in actions_subframe.winfo_children():
            widget.destroy()
        action_frame.actions = []
        
        count = max(int(action_count_var.get()), 1)
        file_options = [f"ファイル{j+1}" for j in range(len(self.file_configs))]
        date_formats = ["%Y/%m/%d", "%Y年%m月%d日", "%Y%m%d", "%Y-%m-%d", "%m/%d/%Y", "%d-%m-%Y", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S"]
        date_options = [datetime.now().strftime(fmt) for fmt in date_formats]
        write_formats = ["文字列", "標準"]
        
        for i in range(count):
            frame = ttk.Frame(actions_subframe)
            frame.pack(fill=tk.X, anchor=tk.NW, pady=2)
            ttk.Label(frame, text=f"実行{i+1}").grid(row=0, column=0, padx=(0, 5))
            source = ttk.Combobox(frame, values=file_options + ["本日の日時", "入力文字"], width=10)
            source.grid(row=0, column=1, padx=(0, 5))
            ttk.Label(frame, text="の").grid(row=0, column=2, padx=(0, 5))
            source_column = ttk.Combobox(frame, width=15)
            source_column.grid(row=0, column=3, padx=(0, 5))
            value_entry = ttk.Entry(frame, width=18)
            ttk.Label(frame, text="を").grid(row=0, column=5, padx=(0, 5))
            dest = ttk.Combobox(frame, values=file_options, width=10)
            dest.grid(row=0, column=6, padx=(0, 5))
            ttk.Label(frame, text="の").grid(row=0, column=7, padx=(0, 5))
            dest_column = ttk.Combobox(frame, width=15)
            dest_column.grid(row=0, column=8, padx=(0, 5))
            ttk.Label(frame, text="に").grid(row=0, column=9, padx=(0, 5))
            method = ttk.Combobox(frame, values=["入力", "末尾に追加", "先頭に追加"], width=10)
            method.grid(row=0, column=10, padx=(0, 5))
            write_format_label = ttk.Label(frame, text="書式:")
            write_format_label.grid(row=0, column=11, padx=(0, 5))
            write_format = ttk.Combobox(frame, values=write_formats, width=10)
            write_format.grid(row=0, column=12, padx=(0, 5))
            source_column_label = ttk.Label(frame, text="")
            source_column_label.grid(row=0, column=3, padx=(0, 5))
            source_column_label.grid_forget()
            
            if self.current_set_name != "新規" and i < len(current_actions):
                act = current_actions[i]
                source.set(act["source"] if act["source"] in file_options + ["本日の日時", "入力文字"] else file_options[0] if file_options else "入力文字")
                dest.set(act["dest"] if act["dest"] in file_options else file_options[0] if file_options else "")
                method.set(act["method"] if act["method"] in ["入力", "末尾に追加", "先頭に追加"] else "入力")
                write_format.set(act["write_format"] if act["write_format"] in write_formats else "文字列")
                if act["source"] == "本日の日時":
                    source_column.configure(values=date_options)
                    source_column.set(act["source_column"] if act["source_column"] in date_options else date_options[0])
                    source_column.grid(row=0, column=3, padx=(0, 5))
                    value_entry.grid_forget()
                elif act["source"].startswith("ファイル"):
                    source_idx = int(act["source"].replace("ファイル", "")) - 1 if act["source"].startswith("ファイル") else 0
                    headers = self.file_configs[source_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                    source_column.configure(values=headers)
                    source_column.set(act["source_column"] if act["source_column"] in headers else headers[0] if headers else "")
                    source_column.grid(row=0, column=3, padx=(0, 5))
                    value_entry.grid_forget()
                else:
                    source_column.grid_forget()
                    value_entry.grid(row=0, column=3, padx=(0, 5))
                    value_entry.delete(0, tk.END)
                    value_entry.insert(0, act["value"])
                dest_idx = int(dest.get().replace("ファイル", "")) - 1 if dest.get().startswith("ファイル") else 0
                headers = self.file_configs[dest_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                dest_column.configure(values=headers)
                dest_column.set(act["dest_column"] if act["dest_column"] in headers else headers[0] if headers else "")
                ext = os.path.splitext(self.file_configs[dest_idx]["path"].get())[1].lower() if dest_idx < len(self.file_configs) else ""
                if ext in ['.xlsx', '.xls']:
                    write_format_label.grid(row=0, column=11, padx=(0, 5))
                    write_format.grid(row=0, column=12, padx=(0, 5))
                else:
                    write_format_label.grid_remove()
                    write_format.grid_remove()
            else:
                source.set(file_options[0] if file_options else "入力文字")
                source_idx = int(source.get().replace("ファイル", "")) - 1 if source.get().startswith("ファイル") else 0
                headers = self.file_configs[source_idx]["headers"] or [f"列{j+1}" for j in range(10)] if source.get().startswith("ファイル") else date_options if source.get() == "本日の日時" else []
                source_column.configure(values=headers)
                source_column.set(headers[0] if headers else date_options[0] if source.get() == "本日の日時" else "")
                if source.get() == "入力文字":
                    source_column.grid_forget()
                    value_entry.grid(row=0, column=3, padx=(0, 5))
                else:
                    source_column.grid(row=0, column=3, padx=(0, 5))
                    value_entry.grid_forget()
                dest.set(file_options[1] if len(file_options) > 1 else file_options[0] if file_options else "")
                dest_idx = int(dest.get().replace("ファイル", "")) - 1 if dest.get().startswith("ファイル") else 0
                headers = self.file_configs[dest_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                dest_column.configure(values=headers)
                dest_column.set(headers[0] if headers else "")
                method.set("入力")
                write_format.set("文字列")
                ext = os.path.splitext(self.file_configs[dest_idx]["path"].get())[1].lower() if dest_idx < len(self.file_configs) else ""
                if ext in ['.xlsx', '.xls']:
                    write_format_label.grid(row=0, column=11, padx=(0, 5))
                    write_format.grid(row=0, column=12, padx=(0, 5))
                else:
                    write_format_label.grid_remove()
                    write_format.grid_remove()
            
            action_frame.actions.append({
                "source": source, "source_column": source_column, "source_column_label": source_column_label,
                "value": value_entry, "dest": dest, "dest_column": dest_column, "method": method,
                "write_format": write_format, "write_format_label": write_format_label
            })
            
            source.bind("<<ComboboxSelected>>", lambda e, s=source, sc=source_column, scl=source_column_label, v=value_entry: self.update_source_column(s, sc, scl, v))
            dest.bind("<<ComboboxSelected>>", lambda e, d=dest, dc=dest_column, wf=write_format, wfl=write_format_label: self.update_dest_column(d, dc, wf, wfl))

    def update_source_column(self, source, source_column, source_column_label, value_entry):
        source_value = source.get()
        date_formats = [
            "%Y/%m/%d", "%Y年%m月%d日", "%Y%m%d", "%Y-%m-%d",
            "%m/%d/%Y", "%d-%m-%Y", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S"
        ]
        date_options = [datetime.now().strftime(fmt) for fmt in date_formats]
        if source_value == "本日の日時":
            source_column.configure(values=date_options)
            current_source_column = source_column.get()
            source_column.set(current_source_column if current_source_column in date_options else date_options[0])
            source_column.grid(row=0, column=3, padx=(0, 5))
            source_column_label.grid_forget()
            value_entry.grid_forget()
        elif source_value.startswith("ファイル"):
            source_idx = int(source_value.replace("ファイル", "")) - 1
            if source_idx < len(self.file_configs):
                headers = self.file_configs[source_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                source_column.configure(values=headers)
                current_source_column = source_column.get()
                source_column.set(current_source_column if current_source_column in headers else headers[0] if headers else "")
                source_column.grid(row=0, column=3, padx=(0, 5))
                source_column_label.grid_forget()
            else:
                source_column.configure(values=[f"列{j+1}" for j in range(10)])
                source_column.set("列1")
                source_column.grid(row=0, column=3, padx=(0, 5))
                source_column_label.grid_forget()
            value_entry.grid_forget()
        else:
            source_column.grid_forget()
            source_column_label.grid_forget()
            value_entry.grid(row=0, column=3, padx=(0, 5))

    def update_dest_column(self, dest, dest_column, write_format, write_format_label):
        dest_idx = int(dest.get().replace("ファイル", "")) - 1 if dest.get().startswith("ファイル") else 0
        if dest_idx < len(self.file_configs):
            headers = self.file_configs[dest_idx]["headers"] or [f"列{j+1}" for j in range(10)]
            dest_column.configure(values=headers)
            current_dest_column = dest_column.get()
            dest_column.set(current_dest_column if current_dest_column in headers else headers[0] if headers else "")
            ext = os.path.splitext(self.file_configs[dest_idx]["path"].get())[1].lower()
            if ext in ['.xlsx', '.xls']:
                write_format_label.grid(row=0, column=11, padx=(0, 5))
                write_format.grid(row=0, column=12, padx=(0, 5))
            else:
                write_format_label.grid_remove()
                write_format.grid_remove()
        else:
            dest_column.configure(values=[f"列{j+1}" for j in range(10)])
            dest_column.set("列1")
            write_format_label.grid_remove()
            write_format.grid_remove()

    def ask_yes_no(self, title, message):
        result = [None]
        dialog = tk.Toplevel(self.root)
        dialog.title(title)
        dialog.transient(self.root)
        dialog.grab_set()
        dialog_width = 300
        dialog_height = 150
        root_x = self.root.winfo_x()
        root_y = self.root.winfo_y()
        root_width = self.root.winfo_width()
        root_height = self.root.winfo_height()
        dialog_x = root_x + (root_width - dialog_width) // 2
        dialog_y = root_y + (root_height - dialog_height) // 2
        dialog.geometry(f"{dialog_width}x{dialog_height}+{dialog_x}+{dialog_y}")
        ttk.Label(dialog, text=message, wraplength=250).pack(pady=20)
        button_frame = ttk.Frame(dialog)
        button_frame.pack(pady=10)
        def set_result(value):
            result[0] = value
            dialog.destroy()
        ttk.Button(button_frame, text="はい", command=lambda: set_result(True)).pack(side=tk.LEFT, padx=10)
        ttk.Button(button_frame, text="いいえ", command=lambda: set_result(False)).pack(side=tk.LEFT, padx=10)
        self.root.wait_window(dialog)
        return result[0]

    def save_current_setting(self):
        if self.current_set_name == "新規":
            set_name = tk.simpledialog.askstring("設定保存", "設定名を入力してください", parent=self.root)
        else:
            set_name = tk.simpledialog.askstring("設定保存", "設定名を入力してください (現在の設定名: {})".format(self.current_set_name), initialvalue=self.current_set_name, parent=self.root)
        if set_name:
            self.saved_settings[set_name] = {
                "file_count": self.file_count_var.get(),
                "set_count": self.set_count_var.get(),
                "file_configs": [
                    {
                        "path": c["path"].get(),
                        "has_header": c["has_header"].get(),
                        "encoding": c["encoding"].get(),
                        "sheet": c["sheet"].get(),
                        "memo": c["memo"].get(),
                        "headers": c["headers"],
                        "read_format": c["read_format"].get()
                    } for c in self.file_configs
                ],
                "sets": [
                    {
                        "conditions": [
                            {
                                "csv1": c["csv1"].get(),
                                "column1": c["column1"].get(),
                                "csv2": c["csv2"].get(),
                                "column2": c["column2"].get(),
                                "operator": c["operator"].get(),
                                "value": c["value"].get(),
                                "logic": c["logic"].get() if c["logic"] else ""
                            } for c in self.set_frames[i].winfo_children()[0].conditions
                        ],
                        "actions": [
                            {
                                "source": a["source"].get(),
                                "source_column": a["source_column"].get() if a["source"].get() != "入力文字" else "",
                                "dest": a["dest"].get(),
                                "dest_column": a["dest_column"].get(),
                                "method": a["method"].get(),
                                "value": a["value"].get() if a["source"].get() == "入力文字" else "",
                                "write_format": a["write_format"].get()
                            } for a in self.set_frames[i].winfo_children()[1].actions
                        ]
                    } for i in range(len(self.set_frames))
                ]
            }
            self.save_settings()
            self.set_name_combo.configure(values=["新規"] + list(self.saved_settings.keys()))
            self.set_name_combo.set(set_name)
            self.current_set_name = set_name

    def delete_current_setting(self):
        if self.current_set_name == "新規":
            if self.ask_yes_no("確認", "現在の設定をクリアしますか？"):
                self.reset_to_default()
        else:
            if self.ask_yes_no("確認", f"設定 '{self.current_set_name}' を削除しますか？"):
                del self.saved_settings[self.current_set_name]
                self.save_settings()
                self.set_name_combo.configure(values=["新規"] + list(self.saved_settings.keys()))
                self.reset_to_default()

    def reset_to_default(self):
        self.current_set_name = "新規"
        self.set_name_combo.set("新規")
        for config in self.file_configs:
            config["path"].set("")
            config["has_header"].set("カラム有")
            config["encoding"].set("自動")
            config["sheet"].set("")
            config["memo"].set("")
            config["read_format"].set("表示値")
            config["headers"] = [f"列{j+1}" for j in range(10)]
            config["sheet_combo"].grid_remove()
            config["read_format_combo"].grid_remove()
        self.file_count_var.set("2")
        self.set_count_var.set("1")
        self.update_file_configs()
        self.update_sets()

    def load_selected_set(self, event):
        set_name = self.set_name_combo.get()
        if set_name == "新規":
            self.reset_to_default()
        else:
            setting = self.saved_settings[set_name]
            self.file_count_var.set(setting["file_count"])
            self.set_count_var.set(setting["set_count"])
            self.update_file_configs()
            
            for i, config in enumerate(setting["file_configs"]):
                self.file_configs[i]["path"].set(config["path"])
                self.file_configs[i]["has_header"].set(config["has_header"])
                self.file_configs[i]["encoding"].set(config["encoding"])
                self.file_configs[i]["sheet"].set(config.get("sheet", ""))
                self.file_configs[i]["memo"].set(config.get("memo", ""))
                self.file_configs[i]["read_format"].set(config.get("read_format", "表示値"))
                self.file_configs[i]["headers"] = config.get("headers", [f"列{j+1}" for j in range(10)])
                if config["path"] and os.path.exists(config["path"]):
                    if os.path.splitext(config["path"])[1].lower() in ['.xlsx', '.xls']:
                        sheet_names = self.get_excel_sheets(config["path"])
                        self.file_configs[i]["sheet_combo"].configure(values=sheet_names)
                        self.file_configs[i]["sheet_combo"].grid() if sheet_names else self.file_configs[i]["sheet_combo"].grid_remove()
                        self.file_configs[i]["read_format_combo"].grid() if sheet_names else self.file_configs[i]["read_format_combo"].grid_remove()
                        if self.file_configs[i]["sheet"].get() not in sheet_names:
                            self.file_configs[i]["sheet"].set(sheet_names[0] if sheet_names else "")
                    else:
                        self.file_configs[i]["sheet_combo"].grid_remove()
                        self.file_configs[i]["read_format_combo"].grid_remove()
                        self.file_configs[i]["sheet"].set("")
                    self.update_file_headers(self.file_configs[i]["path"], i)
                else:
                    self.file_configs[i]["sheet_combo"].grid_remove()
                    self.file_configs[i]["read_format_combo"].grid_remove()
            
            self.update_sets()
            for i, set_config in enumerate(setting["sets"]):
                condition_frame = self.set_frames[i].winfo_children()[0]
                condition_count_var = condition_frame.condition_count_var
                condition_count_var.set(str(max(len(set_config["conditions"]), 1)))
                self.update_conditions(i, condition_count_var)
                for j, cond in enumerate(set_config["conditions"]):
                    if j < len(condition_frame.conditions):
                        condition_frame.conditions[j]["csv1"].set(cond["csv1"])
                        csv1_idx = int(cond["csv1"].replace("ファイル", "")) - 1 if cond["csv1"].startswith("ファイル") else 0
                        headers = self.file_configs[csv1_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                        condition_frame.conditions[j]["column1"].configure(values=headers)
                        condition_frame.conditions[j]["column1"].set(cond["column1"] if cond["column1"] in headers else headers[0] if headers else "")
                        condition_frame.conditions[j]["csv2"].set(cond["csv2"])
                        if cond["csv2"] == "入力文字":
                            condition_frame.conditions[j]["column2"].configure(values=[])
                            condition_frame.conditions[j]["column2"].set("")
                            condition_frame.conditions[j]["value"].delete(0, tk.END)
                            condition_frame.conditions[j]["value"].insert(0, cond["value"])
                        else:
                            csv2_idx = int(cond["csv2"].replace("ファイル", "")) - 1 if cond["csv2"].startswith("ファイル") else 0
                            headers2 = self.file_configs[csv2_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                            condition_frame.conditions[j]["column2"].configure(values=headers2)
                            condition_frame.conditions[j]["column2"].set(cond["column2"] if cond["column2"] in headers2 else headers2[0] if headers2 else "")
                        condition_frame.conditions[j]["operator"].set(cond["operator"])
                        if condition_frame.conditions[j]["logic"]:
                            condition_frame.conditions[j]["logic"].set(cond["logic"])
                        self.update_column2(
                            condition_frame.conditions[j]["csv2"],
                            condition_frame.conditions[j]["value"],
                            condition_frame.conditions[j]["column2"],
                            condition_frame.conditions[j]["no_label"],
                            condition_frame.conditions[j]["to_label"]
                        )
                
                action_frame = self.set_frames[i].winfo_children()[1]
                action_count_var = action_frame.action_count_var
                action_count_var.set(str(max(len(set_config["actions"]), 1)))
                self.update_actions(i, action_count_var)
                for j, act in enumerate(set_config["actions"]):
                    if j < len(action_frame.actions):
                        action_frame.actions[j]["source"].set(act["source"])
                        if act["source"] == "本日の日時":
                            action_frame.actions[j]["source_column"].configure(values=["%Y/%m/%d", "%Y年%m月%d日", "%Y%m%d", "%Y-%m-%d", "%m/%d/%Y", "%d-%m-%Y", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S"])
                            action_frame.actions[j]["source_column"].set(act["source_column"])
                        elif act["source"].startswith("ファイル"):
                            source_idx = int(act["source"].replace("ファイル", "")) - 1 if act["source"].startswith("ファイル") else 0
                            headers = self.file_configs[source_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                            action_frame.actions[j]["source_column"].configure(values=headers)
                            action_frame.actions[j]["source_column"].set(act["source_column"] if act["source_column"] in headers else headers[0] if headers else "")
                        else:
                            action_frame.actions[j]["source_column"].configure(values=[])
                            action_frame.actions[j]["source_column"].set("")
                            action_frame.actions[j]["value"].delete(0, tk.END)
                            action_frame.actions[j]["value"].insert(0, act["value"])
                        action_frame.actions[j]["dest"].set(act["dest"])
                        dest_idx = int(act["dest"].replace("ファイル", "")) - 1 if act["dest"].startswith("ファイル") else 0
                        headers = self.file_configs[dest_idx]["headers"] or [f"列{j+1}" for j in range(10)]
                        action_frame.actions[j]["dest_column"].configure(values=headers)
                        action_frame.actions[j]["dest_column"].set(act["dest_column"] if act["dest_column"] in headers else headers[0] if headers else "")
                        action_frame.actions[j]["method"].set(act["method"])
                        action_frame.actions[j]["write_format"].set(act.get("write_format", "文字列"))
                        self.update_source_column(
                            action_frame.actions[j]["source"],
                            action_frame.actions[j]["source_column"],
                            action_frame.actions[j]["source_column_label"],
                            action_frame.actions[j]["value"]
                        )
                        ext = os.path.splitext(self.file_configs[dest_idx]["path"].get())[1].lower() if dest_idx < len(self.file_configs) else ""
                        if ext in ['.xlsx', '.xls']:
                            action_frame.actions[j]["write_format"].master.grid(row=0, column=12, padx=(0, 5))
                            action_frame.actions[j]["write_format_label"].grid(row=0, column=11, padx=(0, 5))
                        else:
                            action_frame.actions[j]["write_format"].master.grid_remove()
                            action_frame.actions[j]["write_format_label"].grid_remove()
            self.current_set_name = set_name

    def load_file(self):
        self.file_data = []
        self.file_headers = []
        for config in self.file_configs:
            path = config["path"].get()
            if not path:
                messagebox.showerror("エラー", f"ファイル{self.file_configs.index(config)+1}のパスを指定してください")
                return False
            ext = os.path.splitext(path)[1].lower()
            if ext in ['.xlsx', '.xls']:
                try:
                    wb = openpyxl.load_workbook(path, data_only=(config["read_format"].get() == "内部値"))
                    sheet_name = config["sheet"].get() or wb.sheetnames[0]
                    sheet = wb[sheet_name]
                    data = []
                    for row in sheet.iter_rows(values_only=True):
                        row_data = []
                        for cell in row:
                            if cell is None:
                                row_data.append('')
                            elif config["read_format"].get() == "内部値" and isinstance(cell, (int, float, datetime)):
                                row_data.append(cell)
                            else:
                                row_data.append(str(cell))
                        data.append(row_data)
                    if not data or all(not row for row in data):
                        self.file_headers.append([f"列{j+1}" for j in range(10)])
                        self.file_data.append([])
                    else:
                        headers = data[0] if config["has_header"].get() == "カラム有" else [f"列{j+1}" for j in range(len(data[0]))]
                        self.file_headers.append(headers)
                        self.file_data.append(data[1:] if config["has_header"].get() == "カラム有" else data)
                    wb.close()
                except Exception as e:
                    messagebox.showerror("エラー", f"ファイル{self.file_configs.index(config)+1}の読み込みに失敗しました: {str(e)}")
                    return False
            else:  # CSV
                encoding = config["encoding"].get() if config["encoding"].get() != "自動" else self.detect_encoding(path)
                try:
                    delimiter = self.detect_delimiter(path, encoding)
                    with open(path, 'r', encoding=encoding) as f:
                        reader = csv.reader(f, delimiter=delimiter)
                        data = list(reader)
                        if not data or all(not row or all(not cell.strip() for cell in row) for row in data):
                            self.file_headers.append([f"列{j+1}" for j in range(10)])
                            self.file_data.append([])
                        else:
                            headers = data[0] if config["has_header"].get() == "カラム有" else [f"列{j+1}" for j in range(len(data[0]))]
                            self.file_headers.append(headers)
                            self.file_data.append(data[1:] if config["has_header"].get() == "カラム有" else data)
                except Exception as e:
                    messagebox.showerror("エラー", f"ファイル{self.file_configs.index(config)+1}の読み込みに失敗しました: {str(e)}")
                    return False
        return True

    def compare_values(self, val1, val2, operator):
        if operator in ["より大きい", "より小さい"]:
            try:
                if isinstance(val1, (int, float)) and isinstance(val2, (int, float)):
                    num1, num2 = val1, val2
                elif isinstance(val1, datetime) and isinstance(val2, datetime):
                    num1, num2 = val1, val2
                else:
                    num1 = float(re.search(r'\d+\.?\d*', str(val1)).group())
                    num2 = float(re.search(r'\d+\.?\d*', str(val2)).group())
                return (operator == "より大きい" and num1 > num2) or (operator == "より小さい" and num1 < num2)
            except:
                return False
        elif operator == "一致する":
            return str(val1) == str(val2)
        elif operator == "異なる":
            return str(val1) != str(val2)
        elif operator == "含む":
            return str(val2) in str(val1)
        elif operator == "含まない":
            return str(val2) not in str(val1)
        return False

    def execute(self):
        if not self.load_file():
            return

        results = {i: {"actions": {}} for i in range(len(self.set_frames))}
        updates = {i: {} for i in range(len(self.file_configs))}

        for set_idx, set_frame in enumerate(self.set_frames):
            condition_frame = set_frame.winfo_children()[0]
            conditions = condition_frame.conditions
            action_frame = set_frame.winfo_children()[1]
            actions = action_frame.actions

            matches = []
            for csv1_idx, csv1_data in enumerate(self.file_data):
                for row1_idx, row1 in enumerate(csv1_data):
                    match_found = True
                    matched_rows = {i: [] for i in range(len(self.file_data))}
                    matched_rows[csv1_idx].append(row1_idx)

                    for cond_idx, cond in enumerate(conditions):
                        csv1_idx_cond = int(cond["csv1"].get().replace("ファイル", "")) - 1
                        col1_idx = self.file_headers[csv1_idx_cond].index(cond["column1"].get())
                        val1 = self.file_data[csv1_idx_cond][row1_idx][col1_idx] if csv1_idx == csv1_idx_cond else None

                        if cond["csv2"].get() == "入力文字":
                            val2 = cond["value"].get()
                            if val1 is not None and not self.compare_values(val1, val2, cond["operator"].get()):
                                match_found = False
                                break
                        else:
                            csv2_idx = int(cond["csv2"].get().replace("ファイル", "")) - 1
                            col2_idx = self.file_headers[csv2_idx].index(cond["column2"].get())
                            found_match = False
                            for row2_idx, row2 in enumerate(self.file_data[csv2_idx]):
                                val2 = row2[col2_idx]
                                if val1 is not None and self.compare_values(val1, val2, cond["operator"].get()):
                                    matched_rows[csv2_idx].append(row2_idx)
                                    found_match = True
                            if not found_match:
                                match_found = False
                                break

                        if cond_idx < len(conditions) - 1:
                            logic = cond["logic"].get()
                            if logic == "かつ" and not match_found:
                                break
                            elif logic == "または" and match_found:
                                break

                    if match_found:
                        matches.append((csv1_idx, row1_idx, matched_rows))

            results[set_idx]["count"] = len(matches)
            for action_idx, action in enumerate(actions):
                results[set_idx]["actions"][action_idx] = 0
            for csv1_idx, row1_idx, matched_rows in matches:
                for action_idx, action in enumerate(actions):
                    dest_idx = int(action["dest"].get().replace("ファイル", "")) - 1
                    dest_col_idx = self.file_headers[dest_idx].index(action["dest_column"].get())
                    write_format = action["write_format"].get()

                    if dest_idx not in updates:
                        updates[dest_idx] = {}

                    if action["source"].get() == "本日の日時":
                        value = datetime.now().strftime(action["source_column"].get())
                    elif action["source"].get() == "入力文字":
                        value = action["value"].get()
                    else:
                        source_idx = int(action["source"].get().replace("ファイル", "")) - 1
                        source_col_idx = self.file_headers[source_idx].index(action["source_column"].get())
                        source_row_indices = matched_rows.get(source_idx, [row1_idx])
                        value = self.file_data[source_idx][source_row_indices[0]][source_col_idx] if source_row_indices else ""

                    method = action["method"].get()
                    dest_rows = matched_rows.get(dest_idx, [])
                    for dest_row in dest_rows:
                        if dest_row not in updates[dest_idx]:
                            updates[dest_idx][dest_row] = {}
                        if method == "入力":
                            updates[dest_idx][dest_row][dest_col_idx] = value
                        elif method == "末尾に追加":
                            current_value = self.file_data[dest_idx][dest_row][dest_col_idx] if dest_row < len(self.file_data[dest_idx]) else ""
                            updates[dest_idx][dest_row][dest_col_idx] = str(current_value) + str(value)
                        elif method == "先頭に追加":
                            current_value = self.file_data[dest_idx][dest_row][dest_col_idx] if dest_row < len(self.file_data[dest_idx]) else ""
                            updates[dest_idx][dest_row][dest_col_idx] = str(value) + str(current_value)
                        results[set_idx]["actions"][action_idx] += 1

        result_message = ""
        for set_idx in range(len(self.set_frames)):
            for action_idx in range(len(self.set_frames[set_idx].winfo_children()[1].actions)):
                count = results[set_idx]["actions"].get(action_idx, 0)
                result_message += f"セット{set_idx+1}の実行{action_idx+1}を{count}件実行しました。\n"
        if not any(len(updates[i]) > 0 for i in updates):
            result_message = "条件に一致するデータが見つかりませんでした"

        for file_idx, update_data in updates.items():
            path = self.file_configs[file_idx]["path"].get()
            ext = os.path.splitext(path)[1].lower()
            if ext in ['.xlsx', '.xls']:
                try:
                    wb = openpyxl.load_workbook(path)
                    sheet_name = self.file_configs[file_idx]["sheet"].get() or wb.sheetnames[0]
                    sheet_name = self.file_configs[file_idx]["sheet"].get() or wb.sheetnames[0]
                    sheet = wb[sheet_name]
                    has_header = self.file_configs[file_idx]["has_header"].get() == "カラム有"
                    offset = 1 if has_header else 0

                    for row_idx, col_updates in update_data.items():
                        file_row_idx = row_idx + offset
                        for col_idx, value in col_updates.items():
                            write_format = next(
                                (a["write_format"].get() for a in self.set_frames[set_idx].winfo_children()[1].actions
                                 if int(a["dest"].get().replace("ファイル", "")) - 1 == file_idx and
                                 self.file_headers[file_idx].index(a["dest_column"].get()) == col_idx),
                                "文字列"
                            )
                            cell = sheet.cell(row=file_row_idx + 1, column=col_idx + 1)
                            if write_format == "文字列":
                                cell.value = str(value)
                            else:  # 標準
                                try:
                                    if isinstance(value, (int, float)):
                                        cell.value = value
                                    elif isinstance(value, datetime):
                                        cell.value = value
                                    else:
                                        cell.value = str(value)
                                except:
                                    cell.value = str(value)

                    wb.save(path)
                    wb.close()
                except Exception as e:
                    messagebox.showerror("エラー", f"ファイル{file_idx+1}の保存に失敗しました: {str(e)}")
                    return
            else:  # CSV
                try:
                    encoding = self.file_configs[file_idx]["encoding"].get() if self.file_configs[file_idx]["encoding"].get() != "自動" else self.detect_encoding(path)
                    delimiter = self.detect_delimiter(path, encoding)
                    with open(path, 'r', encoding=encoding) as f:
                        reader = csv.reader(f, delimiter=delimiter)
                        data = list(reader)
                        has_header = self.file_configs[file_idx]["has_header"].get() == "カラム有"
                        offset = 1 if has_header else 0

                        for row_idx, col_updates in update_data.items():
                            file_row_idx = row_idx + offset
                            if file_row_idx >= len(data):
                                data.extend([[""] * len(data[0]) for _ in range(file_row_idx - len(data) + 1)])
                            for col_idx, value in col_updates.items():
                                if col_idx >= len(data[file_row_idx]):
                                    data[file_row_idx].extend([""] * (col_idx - len(data[file_row_idx]) + 1))
                                data[file_row_idx][col_idx] = str(value)

                    with open(path, 'w', encoding=encoding, newline='') as f:
                        writer = csv.writer(f, delimiter=delimiter)
                        writer.writerows(data)
                except Exception as e:
                    messagebox.showerror("エラー", f"ファイル{file_idx+1}の保存に失敗しました: {str(e)}")
                    return

        messagebox.showinfo("実行結果", result_message)

if __name__ == "__main__":
    root = tk.Tk()
    app = CSVProcessorApp(root)
    root.mainloop()